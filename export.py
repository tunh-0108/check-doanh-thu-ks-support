# export.py
# ---------------------------------------------------------------------------
# Xuất bảng doanh thu (total_income_by_rooms) ra file Excel (.xlsx):
#   - Cột = các ngày trong tháng, dòng = room_name.
#   - Giá trị ô = doanh thu đêm phòng (tổng nếu ô có nhiều booking).
#   - Màu nền ô = theo Company (OTA) hoặc phương thức thanh toán, do người dùng
#     tự chỉnh ở giao diện trước khi export.
#
# Mỗi ô trong bảng doanh thu lưu ở dạng chuỗi:
#   "[id, doanh_thu, company, pay_type]"                (1 booking hợp lệ)
#   nhiều booking thì nối bằng ", "
#   '["Không tìm thấy chứng từ", id]'                   (booking thiếu chứng từ)
#   '["Thanh toán bằng Thẻ tín dụng", id]'              (booking trả bằng thẻ)
# File này đọc lại các chuỗi đó, tính giá trị + chọn màu.
# ---------------------------------------------------------------------------

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import database


# Các "loại màu" có thể chỉnh (key) và màu mặc định (mã hex, không có dấu #)
COLOR_KEYS = [
    "agoda",
    "traveloka",
    "expedia",
    "ctrip",
    "Chuyển khoản cá nhân",
    "Chuyển khoản công ty",
    "Tiền mặt",
    "Thẻ tín dụng",
    "Không tìm thấy chứng từ",
]

DEFAULT_COLORS = {
    "agoda": "A9C6E8",                  # xanh dương nhạt
    "traveloka": "5B3A1E",              # nâu đậm
    "expedia": "F2B505",                # vàng cam
    "ctrip": "3E7E8C",                  # xanh teal
    "Chuyển khoản cá nhân": "FF00FF",   # hồng cánh sen
    "Chuyển khoản công ty": "D9A9C6",   # hồng nhạt
    "Tiền mặt": "FFFFFF",               # trắng
    "Thẻ tín dụng": "FF0000",           # đỏ
    "Không tìm thấy chứng từ": "D9D9D9",  # xám nhạt
}

OTA_COMPANIES = ("agoda", "traveloka", "expedia", "ctrip")


def _parse_pieces(cell_value):
    """
    Tách một ô bảng doanh thu thành danh sách 'mảnh' (mỗi mảnh là 1 booking).
    Mỗi mảnh là dict có 'kind':
      - 'ota'      : booking hợp lệ qua OTA  -> có 'revenue', 'company'
      - 'walkin'   : booking hợp lệ walk-in/booking -> có 'revenue','company','pay_type'
      - 'credit'   : booking trả bằng thẻ tín dụng (không có doanh thu)
      - 'notfound' : booking thiếu chứng từ (không có doanh thu)
    """
    if cell_value is None:
        return []
    s = str(cell_value).strip()
    if s == "" or s.lower() == "nan":
        return []

    pieces = []
    # Lấy nội dung bên trong từng cặp ngoặc vuông [...]
    for group in re.findall(r"\[([^\]]+)\]", s):
        toks = [t.strip() for t in group.split(",")]
        first = toks[0].strip().strip('"').strip("'")

        if first == "Không tìm thấy chứng từ":
            pieces.append({"kind": "notfound"})
        elif first == "Thanh toán bằng Thẻ tín dụng":
            pieces.append({"kind": "credit"})
        elif len(toks) >= 4:
            company = toks[2]
            kind = "ota" if company in OTA_COMPANIES else "walkin"
            pieces.append({
                "kind": kind,
                "revenue": toks[1],
                "company": company,
                "pay_type": toks[3],
            })
        # nếu không khớp dạng nào thì bỏ qua (an toàn)
    return pieces


def _to_int(text):
    """Chuyển chuỗi doanh thu thành số nguyên; trả về 0 nếu không đọc được."""
    try:
        return int(float(str(text).replace(",", "").strip()))
    except (ValueError, TypeError):
        return 0


def _cell_number(pieces):
    """Tổng doanh thu của các booking hợp lệ trong ô. None nếu không có booking hợp lệ."""
    valids = [p for p in pieces if p["kind"] in ("ota", "walkin")]
    if not valids:
        return None
    return sum(_to_int(p["revenue"]) for p in valids)


def _walk_paytype(piece):
    """
    Quy về 'loại màu' theo phương thức thanh toán cho mảnh walk-in/booking (hoặc thẻ).
    'Chuyển khoản cá nhân + Tiền mặt' dùng chung màu 'Chuyển khoản cá nhân'.
    """
    if piece["kind"] == "credit":
        return "Thẻ tín dụng"
    pt = piece.get("pay_type", "")
    if pt == "Chuyển khoản cá nhân + Tiền mặt":
        return "Chuyển khoản cá nhân"
    return pt


def _pick_color_key(pieces):
    """
    Chọn 'loại màu' cho ô theo thuật toán đã thống nhất. Trả về key trong
    COLOR_KEYS, hoặc None nếu ô để trắng (không có gì).
    """
    otas = [p for p in pieces if p["kind"] == "ota"]
    walks = [p for p in pieces if p["kind"] in ("walkin", "credit")]
    notfounds = [p for p in pieces if p["kind"] == "notfound"]

    if otas or walks:
        if otas and not walks:
            return otas[0]["company"]            # tất cả OTA -> theo booking đầu
        if walks and not otas:
            return _walk_paytype(walks[0])        # tất cả walk-in/booking -> booking đầu
        # lẫn lộn OTA + walk-in/booking -> ưu tiên theo walk-in/booking
        if any(_walk_paytype(p) == "Chuyển khoản công ty" for p in walks):
            return "Chuyển khoản công ty"
        if any(_walk_paytype(p) == "Thẻ tín dụng" for p in walks):
            return "Thẻ tín dụng"
        return _walk_paytype(walks[0])

    if notfounds:
        return "Không tìm thấy chứng từ"
    return None


def _format_day_header(col_name):
    """Đổi tên cột ngày 'D/M' thành 'DD/MM' (vd '1/5' -> '01/05')."""
    parts = str(col_name).split("/")
    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
        return f"{int(parts[0]):02d}/{int(parts[1]):02d}"
    return str(col_name)


def build_revenue_xlsx_bytes(color_settings=None):
    """
    Dựng file Excel doanh thu và trả về dưới dạng bytes (để tải về).
    color_settings: dict {loại màu -> mã hex 'RRGGBB'}. Thiếu loại nào thì dùng mặc định.
    """
    colors = dict(DEFAULT_COLORS)
    if color_settings:
        for k, v in color_settings.items():
            if v:
                colors[k] = str(v).lstrip("#").upper()

    df = database.get_revenue_table_data()
    day_columns = [c for c in df.columns if c != "room_name"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Doanh thu"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill("solid", fgColor="333333")
    header_font = Font(bold=True, color="FFFFFF", name="Arial")

    # --- Hàng tiêu đề ---
    ws.cell(row=1, column=1, value="Room")
    for j, col in enumerate(day_columns, start=2):
        ws.cell(row=1, column=j, value=_format_day_header(col))
    for j in range(1, len(day_columns) + 2):
        c = ws.cell(row=1, column=j)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    # --- Các hàng dữ liệu ---
    for i, (_, row) in enumerate(df.iterrows(), start=2):
        # Cột đầu: tên phòng (hoặc CHI/TỔNG)
        name_cell = ws.cell(row=i, column=1, value=row["room_name"])
        name_cell.font = Font(bold=True, name="Arial")
        name_cell.alignment = center
        name_cell.border = border

        is_special = row["room_name"] in database.NON_ROOM_LABELS

        for j, col in enumerate(day_columns, start=2):
            cell = ws.cell(row=i, column=j)
            cell.border = border
            cell.alignment = center
            cell.font = Font(name="Arial")

            if is_special:
                # CHI / TỔNG: tạm để trống (sẽ định nghĩa logic sau)
                continue

            pieces = _parse_pieces(row[col])
            number = _cell_number(pieces)
            if number is not None:
                cell.value = number

            color_key = _pick_color_key(pieces)
            fill_hex = colors.get(color_key) if color_key else None
            if fill_hex:
                # Thêm 'FF' (alpha đục) để màu hiện đúng trên mọi phần mềm
                cell.fill = PatternFill("solid", fgColor="FF" + fill_hex)

    # Độ rộng cột cho dễ nhìn
    ws.column_dimensions["A"].width = 10
    from openpyxl.utils import get_column_letter
    for j in range(2, len(day_columns) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.freeze_panes = "B2"  # cố định hàng tiêu đề và cột tên phòng

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
